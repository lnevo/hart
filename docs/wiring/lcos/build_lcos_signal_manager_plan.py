#!/usr/bin/env python3
"""Build the LCOS Signal Manager plan workbook from live data.

Every sheet is generated, so nothing depends on an earlier workbook:
  By Signal      ports from docs/wiring/signals/signal_wiring.csv, logic from
                 lcos_signal_logic.csv, watch lists from lcos_signal_plan_notes.csv
  Block Index    BS sensors and their Node:Block comments in tables.xml
  Turnout Index  switches and their Node:Turnout comments in tables.xml
  Legend         how to read the rest
Stops if a signal, block or switch named in the notes is not in the live data.
"""
from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[3]
LCOS = ROOT / "docs/wiring/lcos"
PLAN = LCOS / "lcos_signal_manager_config_plan.xlsx"
WIRING = ROOT / "docs/wiring/signals/signal_wiring.csv"
LOGIC = LCOS / "lcos_signal_logic.csv"
NOTES = LCOS / "lcos_signal_plan_notes.csv"
TABLES = ROOT / "jmri/layouts/hart/output/tables.xml"
PORT_RE = re.compile(r"^C(\d+)-OU(\d+)-(\d+)$")
BLOCK_RE = re.compile(r"Node: (\d+) Block: (\d+)")
TURNOUT_RE = re.compile(r"Node: (\d+) Turnout: (\d+)")

SIGNAL_HEADERS = [
    "Node", "MQTT node", "Location", "Mast", "Head", "LCOS Signal", "UID", "Packed / IH",
    "Port STOP (R)", "Port STOP OU", "Port APPROACH (Y)", "Port APPROACH OU",
    "Port CLEAR (G)", "Port CLEAR OU", "Plant", "Signal role", "Turnout polarity",
    "OR aspect", "Turnouts (name + Node:id)", "Blocks throat A,B (before)",
    "Blocks main C,D (normal)", "Blocks diverge E,F", "Boundary note",
]

LEGEND = [
    "LCOS Signal Manager config plan",
    f"Generated {date.today().isoformat()} by build_lcos_signal_manager_plan.py from live data.",
    "",
    "Sort: Node → LCOS Signal index → Mast/Head",
    "",
    "Ports: flat index for Signal Manager Port A/B/C pickers.",
    "  OU1 pins 1–8 → ports 0–7",
    "  OU2 pins 1–8 → ports 8–15",
    "  OU3 → 16–23, OU4 → 24–31",
    "  Formula: port = (OU - 1) * 8 + (pin - 1)",
    "  Aspect map: Port A = STOP (R), Port B = APPROACH (Y), Port C = CLEAR (G). Field wiring is "
    "Stop, Approach, Clear on consecutive increasing ports (R then Y then G).",
    "  Turnout polarity: Main = Normal (closed); Diverged = Reverse (thrown). "
    "Intermediates 2035/2036 have no plant turnout (—).",
    "  Stock Signal Manager sentence: CLEAR if aligned {Main|Diverged} or set aspect "
    "{Stop|Approach}. Occupied blocks still force STOP. 2-head top = Main/Stop; bottom = "
    "Diverged/Stop (unused disc stays red). Diverge dwarfs = Diverged/Stop. "
    "Balloon 2035/2036 = — / Approach.",
    "  C13: motors on OU2; 8LA/8LB Digicon on OU3. Switch 7/11/13 are TWOSENSOR. "
    "Switch 7/9/11 are on Node 13 OU2.",
    "",
    "Turnouts / Blocks Node:id match JMRI comments:",
    "  Node: N Turnout: T   (T is 0-based LCOS turnout index on that client)",
    "  Node: N Block: B     (B is 1-based in comments; LCOS SM checkbox = B-1)",
    "",
    "Block groups (route-dependent ABS intent):",
    "  Throat A,B = two blocks before the facing-point entrance",
    "  Main C,D   = two blocks on the normal / MAIN path",
    "  Diverge E,F = two blocks on the DIVERGED path",
    "  When a leg has only one detector, only one entry is listed.",
    "",
    "Cross-node: if Node:id is not this signal client, subscribe that node and use Remote 8–15",
    "  (remote checkbox = 8 + neighbor local idx). Confirm peer order in LCOS object listen list.",
    "",
    "Sources: docs/wiring/signals/signal_wiring.csv, docs/wiring/lcos/lcos_signal_logic.csv,",
    "  docs/wiring/lcos/lcos_signal_plan_notes.csv, jmri/layouts/hart/output/tables.xml",
    "",
    "Plant block assignments are the intended ABS watch lists for native LCOS / proposed route UI.",
    "Stock Signal Manager cannot gate main vs diverge blocks on one head — prefer one head per "
    "entrance where route-dependent logic matters.",
]


def port_index(port_id: str) -> int:
    m = PORT_RE.match(port_id)
    if not m:
        raise SystemExit(f"bad port_id {port_id}")
    return (int(m.group(2)) - 1) * 8 + (int(m.group(3)) - 1)


def load_heads() -> dict[tuple[str, str], dict]:
    lamps: dict[tuple[str, int], dict[str, str]] = defaultdict(dict)
    meta: dict[tuple[str, int], dict[str, str]] = {}
    with WIRING.open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            key = (row["mast_user_name"], int(row["packed"]))
            lamps[key][row["lamp_color"]] = row["port_id"]
            meta[key] = row
    return {
        (mast, meta[(mast, packed)]["disc_role"]): {**meta[(mast, packed)], "lamps": found}
        for (mast, packed), found in lamps.items()
    }


def load_keyed(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as f:
        return {(row["mast"], row["head"]): row for row in csv.DictReader(f)}


def load_live_ids() -> tuple[dict[str, tuple[int, int]], dict[str, tuple[int, int]]]:
    """Live BS name → (node, block) and Switch name → (node, turnout)."""
    root = ET.parse(TABLES).getroot()
    blocks: dict[str, tuple[int, int]] = {}
    for sensor in root.iter("sensor"):
        name = (sensor.findtext("userName") or "").strip()
        m = BLOCK_RE.search(sensor.findtext("comment") or "")
        if name.startswith("BS ") and m:
            blocks[name] = (int(m.group(1)), int(m.group(2)))
    turnouts: dict[str, tuple[int, int]] = {}
    for turnout in root.iter("turnout"):
        name = (turnout.findtext("userName") or "").strip()
        m = TURNOUT_RE.search(turnout.findtext("comment") or "")
        if name.startswith("Switch ") and m:
            turnouts[name] = (int(m.group(1)), int(m.group(2)))
    return blocks, turnouts


def describe(names: str, live: dict, label: str, unknown: set[str]) -> str:
    parts = []
    for name in filter(None, (p.strip() for p in names.split("|"))):
        if name not in live:
            unknown.add(name)
            parts.append(name)
            continue
        node, ident = live[name]
        index = ident - 1 if label == "Block" else ident
        parts.append(f"{name} (Node: {node} {label}: {ident}; LCOS idx {index})")
    return " | ".join(parts)


def bold_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)


def signal_rows(blocks, turnouts) -> tuple[list[list], set[str]]:
    heads = load_heads()
    logic = load_keyed(LOGIC)
    notes = load_keyed(NOTES)
    missing = (set(heads) ^ set(logic)) | (set(heads) ^ set(notes))
    if missing:
        raise SystemExit(f"signals differ between wiring, logic and notes: {sorted(missing)}")
    unknown: set[str] = set()
    rows = []
    for key, head in heads.items():
        ports = head["lamps"]
        note, rule = notes[key], logic[key]
        rows.append([
            head["parent_node_id"], int(head["mqtt_node"]), head["board_location"], key[0], key[1],
            int(head["signal_index"]), int(head["uid"]), head["packed"],
            port_index(ports["R"]), ports["R"], port_index(ports["Y"]), ports["Y"],
            port_index(ports["G"]), ports["G"], note["plant"], note["signal_role"],
            rule["turnout_polarity"], rule["or_aspect"],
            describe(note["turnouts"], turnouts, "Turnout", unknown),
            describe(note["blocks_throat"], blocks, "Block", unknown),
            describe(note["blocks_main"], blocks, "Block", unknown),
            describe(note["blocks_diverge"], blocks, "Block", unknown),
            note["boundary_note"],
        ])
    rows.sort(key=lambda r: (int(r[0][1:]), r[5], r[3], r[4]))
    return rows, unknown


def main() -> int:
    blocks, turnouts = load_live_ids()
    rows, unknown = signal_rows(blocks, turnouts)
    if unknown:
        raise SystemExit(f"names not in live tables: {sorted(unknown)}")

    wb = Workbook()
    legend = wb.active
    legend.title = "Legend"
    for line in LEGEND:
        legend.append([line])
    legend["A1"].font = Font(bold=True)

    signals = wb.create_sheet("By Signal")
    signals.append(SIGNAL_HEADERS)
    for row in rows:
        signals.append(row)
    bold_header(signals)

    block_sheet = wb.create_sheet("Block Index")
    block_sheet.append(["Node", "LCOS idx (0-7)", "JMRI Block #", "Node:id", "BS name"])
    for name, (node, block) in sorted(blocks.items(), key=lambda kv: kv[1]):
        block_sheet.append([f"C{node}", block - 1, block, f"Node: {node} Block: {block}", name])
    bold_header(block_sheet)

    turnout_sheet = wb.create_sheet("Turnout Index")
    turnout_sheet.append(["Node", "LCOS idx", "Node:id", "Switch name"])
    for name, (node, turnout) in sorted(turnouts.items(), key=lambda kv: kv[1]):
        turnout_sheet.append([f"C{node}", turnout, f"Node: {node} Turnout: {turnout}", name])
    bold_header(turnout_sheet)

    wb.save(PLAN)
    print(f"wrote {PLAN.relative_to(ROOT)}: {len(rows)} heads, {len(blocks)} blocks, {len(turnouts)} turnouts")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
