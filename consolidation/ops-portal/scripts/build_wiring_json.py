#!/usr/bin/env python3
"""Build the public wiring viewer data (behind/wiring.html) from the LCOS workbooks.

Reads the LCOS Layout Inventory and the signal-split workbook and writes:

  data/wiring-inventory.json  nodes with DNOU8 / DNIN8 / block-detector port maps,
                              signal heads with their R/Y/G pins, turnouts, block sensors
  data/signal-splits.json     per control point: each signal's routes, the turnout
                              positions and occupancy behind each aspect, and the lamp pins

Re-run after a version bump:

  python3 consolidation/ops-portal/scripts/build_wiring_json.py \
      --inventory docs/wiring/LCOS_Layout_Inventory_v86.xlsx \
      --splits docs/wiring/signals_split_v9.xlsx

With no arguments it picks the highest-numbered workbook of each kind in docs/wiring/
(the live git copy). Data problems found along the way are printed as warnings; they do
not stop the build. Network details (IP addresses, hostnames, serial ports, machine paths,
credentials) are scrubbed from every cell before it reaches the JSON.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path

import openpyxl

PORTAL = Path(__file__).resolve().parents[1]
ROOT = PORTAL.parents[1]
WIRING = ROOT / "docs" / "wiring"
OUT_INV = PORTAL / "data" / "wiring-inventory.json"
OUT_SPLIT = PORTAL / "data" / "signal-splits.json"

WARNINGS: list[str] = []
SCRUBBED: list[str] = []


def warn(msg: str) -> None:
    WARNINGS.append(msg)


# —— Cell cleanup ————————————————————————————————————————————————

# Anything that looks like a network detail, a machine path or a secret. A matching cell
# is dropped whole rather than half-redacted.
UNSAFE = [
    re.compile(r"\b\d{1,3}(?:\.\d{1,3}){3}\b"),  # IPv4
    re.compile(r"\b(?:[0-9a-f]{2}:){5}[0-9a-f]{2}\b", re.I),  # MAC
    re.compile(r"\b[\w-]+\.(?:local|lan|home|internal)\b", re.I),  # mDNS / LAN hosts
    re.compile(r"\b(?:mqtt|tcp|ws|wss|http|https)://", re.I),
    re.compile(r"\bCOM\d+\b"),
    re.compile(r"/dev/(?:tty|cu)\S*"),
    re.compile(r"(?:/Users/|/home/|[A-Z]:\\)"),
    re.compile(r"\b(?:password|passwd|pwd|secret|token|api[_ -]?key|ssid|wi-?fi|psk|credential)s?\b", re.I),
    re.compile(r"\bbroker\b", re.I),
]


def clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.isoformat()[:10]
    s = re.sub(r"\s+", " ", str(v)).strip()
    for pat in UNSAFE:
        if pat.search(s):
            SCRUBBED.append(s)
            return ""
    return s


def sheet_rows(ws) -> list[list[str]]:
    """Every row as cleaned strings, empty rows dropped."""
    out = []
    for row in ws.iter_rows(values_only=True):
        cells = [clean(c) for c in row]
        if any(cells):
            out.append(cells)
    return out


def table(rows: list[list[str]], header_idx: int, stop=None) -> list[dict]:
    """Rows under header_idx as dicts keyed by header, trailing empty columns dropped."""
    header = rows[header_idx]
    width = max((i + 1 for i, h in enumerate(header) if h), default=0)
    header = header[:width]
    out = []
    for r in rows[header_idx + 1 :]:
        if stop and stop(r):
            break
        d = {header[i]: (r[i] if i < len(r) else "") for i in range(width) if header[i]}
        if any(d.values()):
            out.append(d)
    return out


def strip_history(note: str) -> str:
    """Drop rename history ("was S3-6 G", "renamed for linear6") that only confuses a reader.

    The inventory's "was" notes disagree with each other (DNOU8 and DigiconSignals give
    different old IDs for the same pin), so none of them are published.
    """
    note = re.sub(r"\s*\([^)]*\bwas\b[^)]*\)", "", note)
    parts = re.split(r"\s*[;|]\s*", note)
    keep = [p for p in parts if p and not re.match(r"(?i)was\b", p) and "renamed for" not in p.lower()]
    return "; ".join(keep).strip()


def version_of(path: Path) -> str:
    m = re.search(r"_v(\d+)", path.stem)
    return f"v{m.group(1)}" if m else path.stem


def latest(pattern: str) -> Path | None:
    found = sorted(WIRING.glob(pattern), key=lambda p: int(re.search(r"_v(\d+)", p.stem).group(1)))
    return found[-1] if found else None


def split_list(s: str) -> list[str]:
    return [x.strip() for x in s.split(",") if x.strip()] if s else []


PORT_RE = re.compile(r"^(?P<node>[A-Z]+\d+)-(?P<board>(?:OU|IN|B)\d+)-(?P<pin>\d+)$")
HEAD_RE = re.compile(r"^Head (?P<head>.+?) (?P<color>[RYG])$")


def parse_port(pid: str):
    m = PORT_RE.match(pid)
    if not m:
        return None
    return m.group("node"), m.group("board"), int(m.group("pin"))


def board_sort(b: str):
    m = re.match(r"([A-Z]+)(\d+)", b)
    return (m.group(1), int(m.group(2))) if m else (b, 0)


def node_sort(n: str):
    m = re.match(r"([A-Z]+)(\d+)", n)
    return ({"M": 0, "C": 1, "D": 2}.get(m.group(1), 3), int(m.group(2))) if m else (9, 0)


# —— Inventory ——————————————————————————————————————————————————

def out_kind(device: str, dtype: str) -> tuple[str, str]:
    """(kind, lamp color) for a DNOU8 pin."""
    m = HEAD_RE.match(device)
    if m:
        return "lamp", m.group("color")
    t = dtype.lower()
    if "turnout" in t:
        return "motor", ""
    if "cal" in t:
        return "cal", ""
    if "relay" in t:
        return "relay", ""
    if "reserved" in t or device.lower() == "reserved":
        return "reserved", ""
    if not device:
        return "spare", ""
    return "other", ""


def in_kind(device: str, dtype: str) -> str:
    t = dtype.lower()
    if "feedback" in t:
        return "feedback"
    if "button" in t:
        return "button"
    if not device:
        return "spare"
    return "other"


def pick_location(locs: list[str]) -> str:
    """The board's place: the most common specific location among its used pins."""
    locs = [l for l in locs if l]
    if not locs:
        return ""
    specific = [l for l in locs if "(" in l]
    return Counter(specific or locs).most_common(1)[0][0]


def build_inventory(path: Path) -> tuple[dict, dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {ws.title: sheet_rows(ws) for ws in wb.worksheets}
    for need in ("Nodes", "DNOU8", "DNIN8", "TurnoutSummary", "BlockSensors", "DigiconSignals"):
        if need not in sheets:
            raise SystemExit(f"{path.name}: missing sheet {need}")

    # Nodes
    node_rows = table(sheets["Nodes"], 0)
    nodes: "OrderedDict[str, dict]" = OrderedDict()
    for r in node_rows:
        nid = r.get("Node ID", "")
        if not nid:
            continue
        nodes[nid] = {
            "id": nid,
            "legacy": r.get("Legacy Node ID", "") if r.get("Legacy Node ID") != nid else "",
            "address": r.get("Address", ""),
            "location": r.get("Location", ""),
            "type": r.get("Type", ""),
            "turnouts": split_list(r.get("Turnout Names", "")),
            "signals": split_list(r.get("Signal Names", "")),
            "blockNames": split_list(r.get("Block Section Names", "")),
            "summary": {
                "outputs12v": {"used": r.get("12V Used", ""), "free": r.get("12V Free", "")},
                "outputs5v": {"used": r.get("5V Used", ""), "free": r.get("5V Free", "")},
                "inputs": {"used": r.get("Input Used", ""), "free": r.get("Input Free", "")},
            },
            "outputs": [],
            "inputs": [],
            "detectors": [],
            "declared": {
                "DNOU8": split_list(r.get("DNOU8", "")),
                "DNIN8": split_list(r.get("DNIN8", "")),
                "BS": split_list(r.get("BS Devices", "")),
            },
        }

    def node_for(nid: str, sheet: str, pid: str):
        if nid not in nodes:
            warn(f"{sheet}: {pid} names node {nid}, which is not on the Nodes sheet")
            nodes[nid] = {"id": nid, "legacy": "", "address": "", "location": "", "type": "",
                          "turnouts": [], "signals": [], "blockNames": [], "summary": {},
                          "outputs": [], "inputs": [], "detectors": [], "declared": {}}
        return nodes[nid]

    # DNOU8
    boards: dict[tuple[str, str], dict] = {}
    seen_ports: set[str] = set()
    device_ports: dict[str, list[str]] = defaultdict(list)
    out_pins: dict[str, dict] = {}
    for r in table(sheets["DNOU8"], 0):
        pid = r.get("Output Port ID", "")
        p = parse_port(pid)
        if not p:
            warn(f"DNOU8: unreadable port id {pid!r}")
            continue
        nid, board, pin = p
        if r.get("Parent Node ID") and r["Parent Node ID"] != nid:
            warn(f"DNOU8: {pid} has Parent Node {r['Parent Node ID']}")
        if pid in seen_ports:
            warn(f"DNOU8: port {pid} listed twice")
            continue
        seen_ports.add(pid)
        device = r.get("Connected Device", "")
        kind, color = out_kind(device, r.get("Device Type", ""))
        b = boards.setdefault((nid, board), {"board": board, "rail": Counter(), "pins": [], "_locs": []})
        if r.get("Voltage Rail (12V/5V)"):
            b["rail"][r["Voltage Rail (12V/5V)"]] += 1
        rec = {"pin": pin, "port": pid, "device": device, "kind": kind}
        if color:
            rec["color"] = color
            rec["head"] = HEAD_RE.match(device).group("head")
        if r.get("Board Location"):
            rec["place"] = r["Board Location"]
            if kind not in ("spare", "reserved", "cal"):
                b["_locs"].append(r["Board Location"])
        if r.get("Notes"):
            if kind == "lamp":
                rec["_note"] = r["Notes"]  # route info is lifted onto the head; not published per pin
            elif strip_history(r["Notes"]):
                rec["note"] = strip_history(r["Notes"])
        b["pins"].append(rec)
        out_pins[pid] = rec
        if device and kind not in ("spare", "reserved", "cal"):
            device_ports[device].append(pid)
    for dev, ports in device_ports.items():
        if len(ports) > 1:
            warn(f"DNOU8: {dev!r} is on more than one pin: {', '.join(ports)}")
    for (nid, board), b in sorted(boards.items(), key=lambda kv: (node_sort(kv[0][0]), board_sort(kv[0][1]))):
        rails = list(b["rail"])
        if len(rails) > 1:
            warn(f"DNOU8: {nid}-{board} mixes rails {rails}")
        pins = sorted(b["pins"], key=lambda x: x["pin"])
        have = {x["pin"] for x in pins}
        missing = [i for i in range(1, 9) if i not in have]
        if missing:
            warn(f"DNOU8: {nid}-{board} has no row for pin(s) {missing}")
        node_for(nid, "DNOU8", f"{nid}-{board}")["outputs"].append({
            "board": board,
            "rail": b["rail"].most_common(1)[0][0] if b["rail"] else "",
            "place": pick_location(b["_locs"]),
            "pins": pins,
        })

    # DNIN8
    in_boards: dict[tuple[str, str], dict] = {}
    seen_in: set[str] = set()
    in_device_ports: dict[str, list[str]] = defaultdict(list)
    for r in table(sheets["DNIN8"], 0):
        pid = r.get("Input Port ID", "")
        p = parse_port(pid)
        if not p:
            warn(f"DNIN8: unreadable port id {pid!r}")
            continue
        nid, board, pin = p
        if pid in seen_in:
            warn(f"DNIN8: port {pid} listed twice")
            continue
        seen_in.add(pid)
        device = r.get("Connected Device", "")
        kind = in_kind(device, r.get("Device Type (Sensor/Switch/etc.)", ""))
        rec = {"pin": pin, "port": pid, "device": device, "kind": kind}
        if r.get("Board Location"):
            rec["place"] = r["Board Location"]
        if strip_history(r.get("Notes", "")):
            rec["note"] = strip_history(r["Notes"])
        in_boards.setdefault((nid, board), {"board": board, "pins": [], "_locs": []})
        in_boards[(nid, board)]["pins"].append(rec)
        if device:
            in_boards[(nid, board)]["_locs"].append(r.get("Board Location", ""))
            in_device_ports[device].append(pid)
    for dev, ports in in_device_ports.items():
        if len(ports) > 1:
            warn(f"DNIN8: {dev!r} is on more than one input: {', '.join(ports)}")
    for (nid, board), b in sorted(in_boards.items(), key=lambda kv: (node_sort(kv[0][0]), board_sort(kv[0][1]))):
        pins = sorted(b["pins"], key=lambda x: x["pin"])
        missing = [i for i in range(1, 9) if i not in {x["pin"] for x in pins}]
        if missing:
            warn(f"DNIN8: {nid}-{board} has no row for pin(s) {missing}")
        node_for(nid, "DNIN8", f"{nid}-{board}")["inputs"].append(
            {"board": board, "place": pick_location(b["_locs"]), "pins": pins})

    # Block sensors
    blocks = []
    det: dict[tuple[str, str], list] = defaultdict(list)
    seen_bs: set[str] = set()
    for r in table(sheets["BlockSensors"], 0):
        pid = r.get("Port ID", "")
        nid = r.get("Node ID", "") or (pid.split("-")[0] if pid else "")
        if pid in seen_bs:
            warn(f"BlockSensors: {pid} listed twice")
            continue
        seen_bs.add(pid)
        rec = {
            "port": pid,
            "node": nid,
            "detector": r.get("B Device ID", ""),
            "channel": r.get("B Number (1-2)", ""),
            "block": r.get("Block Section Name", ""),
            "track": r.get("Connected Track Segment", ""),
            "area": r.get("Block Location", ""),
            "district": r.get("Power District", ""),
            "note": strip_history(r.get("Notes", "")),
        }
        if not rec["block"]:
            warn(f"BlockSensors: {pid} has no block name")
        blocks.append(rec)
        det[(nid, rec["detector"])].append(rec)
    for (nid, d), chans in sorted(det.items(), key=lambda kv: (node_sort(kv[0][0]), kv[0][1])):
        node_for(nid, "BlockSensors", d)["detectors"].append({
            "detector": d,
            "channels": [{"channel": c["channel"], "port": c["port"], "block": c["block"]} for c in
                         sorted(chans, key=lambda c: str(c["channel"]))],
        })

    # Signal heads: Digicon searchlights (DigiconSignals, lower deck) + RGB LED heads (DNOU8, upper deck)
    dig_rows = sheets["DigiconSignals"]
    hdr = next(i for i, r in enumerate(dig_rows) if r and r[0] == "port_id")
    mast_hdr = next((i for i, r in enumerate(dig_rows) if r and r[0] == "mast_user_name"), None)
    port_rows = table(dig_rows[: (mast_hdr - 1) if mast_hdr else None], hdr)
    mast_rows = table(dig_rows, mast_hdr) if mast_hdr else []

    heads: "OrderedDict[str, dict]" = OrderedDict()
    for r in port_rows:
        name = r.get("user_name", "")
        m = HEAD_RE.match(name)
        if not m:
            continue
        head = m.group("head")
        color = r.get("lamp_color") or m.group("color")
        mast = r.get("mast_user_name", "").replace("Mast ", "")
        h = heads.setdefault(head, {
            "head": head, "mast": mast, "disc": r.get("disc_role", ""),
            "node": r.get("parent_node_id", ""), "place": r.get("board_location", ""),
            "packed": r.get("system_name", ""), "type": "Digicon searchlight", "lamps": {},
        })
        if color in h["lamps"]:
            warn(f"DigiconSignals: {head} lamp {color} on both {h['lamps'][color]} and {r['port_id']}")
        h["lamps"][color] = r.get("port_id", "")
        dn = out_pins.get(r.get("port_id", ""))
        if not dn:
            warn(f"DigiconSignals: {head} {color} on {r.get('port_id')} which is not on the DNOU8 sheet")
        elif dn.get("head") != head or dn.get("color") != color:
            warn(f"DigiconSignals: {head} {color} on {r['port_id']}, but DNOU8 has {dn.get('device') or 'nothing'} there")

    for pid, rec in sorted(out_pins.items(), key=lambda kv: (node_sort(kv[0].split("-")[0]), kv[0])):
        if rec["kind"] != "lamp":
            continue
        head = rec["head"]
        if head in heads and heads[head]["type"] == "Digicon searchlight":
            if heads[head]["lamps"].get(rec["color"]) != pid:
                warn(f"DNOU8: {rec['device']} on {pid} is not in DigiconSignals")
            continue
        note = rec.get("_note", "")
        parts = [p.strip() for p in note.split("|")]
        mast = ""
        mm = re.search(r"mast (Switch \d+[^|]*)", note)
        if mm:
            mast = mm.group(1).strip()
        h = heads.setdefault(head, {
            "head": head, "mast": mast, "disc": "", "node": pid.split("-")[0],
            "place": rec.get("place", ""), "packed": "", "type": "RGB LED", "lamps": {},
            "cp": parts[0] if parts and re.match(r"CP\d", parts[0]) else "",
            "route": parts[1] if len(parts) > 1 and "→" in parts[1] else "",
        })
        if rec["color"] in h["lamps"]:
            warn(f"DNOU8: {head} lamp {rec['color']} on both {h['lamps'][rec['color']]} and {pid}")
        h["lamps"][rec["color"]] = pid
    for h in heads.values():
        miss = [c for c in "RYG" if c not in h["lamps"]]
        if miss:
            warn(f"Signals: head {h['head']} has no {'/'.join(miss)} pin")
        nodes_used = sorted({p.split("-")[0] for p in h["lamps"].values()})
        if len(nodes_used) > 1:
            warn(f"Signals: head {h['head']} spans nodes {nodes_used}")

    masts = []
    for r in mast_rows:
        masts.append({
            "mast": r.get("mast_user_name", "").replace("Mast ", ""),
            "node": r.get("parent_node_id", ""),
            "heads": r.get("heads", ""),
            "appearance": r.get("appearance", ""),
            "ports": r.get("port_ids", "").split(),
        })
    mast_by = {m["mast"]: m for m in masts}
    for h in heads.values():
        m = mast_by.get(h["mast"])
        if m:
            h["appearance"] = m["appearance"]

    # Node signal names vs the heads actually wired there
    for n in nodes.values():
        wired = {h["head"].split(" ")[0] for h in heads.values() if h["node"] == n["id"]}
        listed = set(n["signals"])
        extra = sorted(listed - wired)
        if extra and n["outputs"]:
            warn(f"Nodes: {n['id']} lists signal(s) {extra} with no lamp pins on that node")

    # Turnouts
    turnouts = []
    for r in table(sheets["TurnoutSummary"], 0):
        name = r.get("Turnout", "")
        if not name:
            continue
        motor = {pos: device_ports.get(f"{name} {pos}", []) for pos in ("N", "R")}
        if not motor["N"] and not motor["R"]:
            warn(f"TurnoutSummary: {name} has no motor pins on the DNOU8 sheet")
        t = {
            "turnout": name,
            "node": r.get("Parent Node", ""),
            "area": r.get("Location/Area", ""),
            "elements": r.get("Switch Elements", ""),
            "motorN": ", ".join(motor["N"]),
            "motorR": ", ".join(motor["R"]),
            "button": r.get("Button", ""),
            "fbN": r.get("Feedback N", ""),
            "fbR": r.get("Feedback R", ""),
            "signals": {},
        }
        motor_nodes = {p.split("-")[0] for p in motor["N"] + motor["R"]}
        if motor_nodes and t["node"] and t["node"] not in motor_nodes:
            warn(f"TurnoutSummary: {name} parent node {t['node']}, but its motor is on {sorted(motor_nodes)}")
        for key, label in (("entry", "Entry Signal"), ("normal", "Normal Exit Signal"), ("reverse", "Reverse Exit Signal")):
            sig = r.get(label, "")
            base = label.replace(" Signal", "")
            ports = {c: r.get(f"{base} Signal {c} Port", "") or r.get(f"{base} {c} Port", "") for c in "RYG"}
            if sig:
                t["signals"][key] = sig
            # Cross-check the sheet's lamp pins against the DNOU8 pin assignments
            for c, pid in ports.items():
                if not pid:
                    continue
                dn = out_pins.get(pid)
                who = dn.get("device") if dn else None
                if not sig:
                    warn(f"TurnoutSummary: {name} {base.lower()} has {c} pin {pid} but no signal name")
                elif not dn:
                    warn(f"TurnoutSummary: {name} {base.lower()} {sig} {c} → {pid}, not on the DNOU8 sheet")
                elif dn.get("head", "").split(" ")[0] != sig.split(" ")[0] or dn.get("color") != c:
                    warn(f"TurnoutSummary: {name} {base.lower()} {sig} {c} → {pid}, but DNOU8 has {who or 'nothing (spare)'} there")
        for f, want in (("button", "BTN"), ("fbN", "N FB"), ("fbR", "R FB")):
            pid = t[f]
            if not pid:
                continue
            dev = next((x["device"] for b in nodes.get(pid.split("-")[0], {}).get("inputs", []) for x in b["pins"] if x["port"] == pid), None)
            if dev != f"{name} {want}":
                warn(f"TurnoutSummary: {name} {f} → {pid}, but DNIN8 has {dev or 'nothing'} there")
        turnouts.append(t)

    # Declared boards vs boards found
    for n in nodes.values():
        dec = n.pop("declared", {})
        found_ou = [b["board"] for b in n["outputs"]]
        found_in = [b["board"] for b in n["inputs"]]
        if dec.get("DNOU8") is not None and sorted(dec.get("DNOU8", [])) != sorted(found_ou):
            warn(f"Nodes: {n['id']} declares DNOU8 {dec.get('DNOU8')} but the DNOU8 sheet has {found_ou}")
        if dec.get("DNIN8") is not None and sorted(dec.get("DNIN8", [])) != sorted(found_in):
            warn(f"Nodes: {n['id']} declares DNIN8 {dec.get('DNIN8')} but the DNIN8 sheet has {found_in}")

    for b in boards.values():
        b.pop("_locs", None)
    for rec in out_pins.values():
        rec.pop("_note", None)
    modified = wb.properties.modified.date().isoformat() if wb.properties.modified else ""
    inv = {
        "source": {"file": path.name, "version": version_of(path), "modified": modified},
        "nodes": sorted(nodes.values(), key=lambda n: node_sort(n["id"])),
        "signals": sorted(heads.values(), key=lambda h: (node_sort(h["node"]), h["head"])),
        "masts": masts,
        "turnouts": turnouts,
        "blocks": blocks,
    }
    return inv, {"out_pins": out_pins, "heads": heads}


# —— Signal splits ——————————————————————————————————————————————————

def build_splits(path: Path, ctx: dict) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    heads = ctx["heads"]
    out_pins = ctx["out_pins"]

    def lamp_pins(sig: str):
        h = heads.get(sig)
        if h:
            return h["head"], h["lamps"], h["node"]
        return None, None, None

    cps = []
    for ws in wb.worksheets:
        m = re.match(r"control_point_(\d+)$", ws.title)
        if not m:
            continue
        rows = sheet_rows(ws)
        if not rows:
            continue
        header = rows[0]
        width = max(i + 1 for i, h in enumerate(header) if h)
        header = header[:width]
        idx = {h: i for i, h in enumerate(header) if h}
        sw_cols = [(i, h) for i, h in enumerate(header) if h.startswith("Switch ")]
        block_cols = [i for i, h in enumerate(header) if re.match(r"Block ?\d$", h)]
        get = lambda r, k: r[idx[k]] if k in idx and idx[k] < len(r) else ""
        # A switch heading used twice is a double crossover's two halves (Switch 23 → 23a, 23b).
        dups = {h for h, c in Counter(h for _, h in sw_cols).items() if c > 1}
        if dups:
            warn(f"{ws.title}: switch column(s) {sorted(dups)} appear twice; shown as a / b halves")
            seen: Counter = Counter()
            renamed = []
            for i, h in sw_cols:
                if h in dups:
                    renamed.append((i, h + "ab"[seen[h]]))
                    seen[h] += 1
                else:
                    renamed.append((i, h))
            sw_cols = renamed

        area = ""
        signals: "OrderedDict[str, dict]" = OrderedDict()
        for r in rows[1:]:
            sig = get(r, "Signal")
            if not sig:
                continue
            area = area or get(r, "Location")
            s = signals.get(sig)
            if s is None:
                head, lamps, node = lamp_pins(sig)
                s = {"signal": sig, "mast": get(r, "Mast"), "routes": []}
                if head:
                    s["head"] = head
                    s["node"] = node
                    s["lamps"] = lamps
                elif re.match(r"S\d-\d+$", sig):
                    s["planned"] = True  # Nov 2025 RGB ID; the lower deck now has Digicon heads
                else:
                    warn(f"{ws.title}: signal {sig} has no lamp pins in the inventory")
                signals[sig] = s
            sw = OrderedDict()
            for i, h in sw_cols:
                v = r[i] if i < len(r) else ""
                if v:
                    if h in sw and sw[h] != v:
                        warn(f"{ws.title}: {sig} row sets {h} both {sw[h]} and {v}")
                    sw[h] = v
            blocks_ = [r[i] for i in block_cols if i < len(r) and r[i]]
            key = (get(r, "Route_From"), get(r, "Route_To"), tuple(blocks_), tuple(sw.items()))
            route = next((x for x in s["routes"] if x["_key"] == key), None)
            if route is None:
                route = {
                    "_key": key,
                    "from": get(r, "Route_From") or get(r, "Signal_Block"),
                    "to": get(r, "Route_To"),
                    "blocks": blocks_,
                    "switches": [{"switch": k, "pos": v} for k, v in sw.items()],
                    "aspects": [],
                }
                s["routes"].append(route)
            b1, b2 = get(r, "B1").upper() == "X", get(r, "B2").upper() == "X"
            when = "B1" if b1 else ("B2" if b2 else "clear")
            if b1 and b2:
                warn(f"{ws.title}: {sig} row marks both B1 and B2 occupied")
            aspect = get(r, "BO_Aspect").strip()
            rule = get(r, "BO_Rule")
            if not aspect:
                warn(f"{ws.title}: {sig} {route['from']}→{route['to']} has an empty aspect")
            if any(a["when"] == when for a in route["aspects"]):
                warn(f"{ws.title}: {sig} {route['from']}→{route['to']} lists the '{when}' case twice")
            route["aspects"].append({"when": when, "aspect": aspect, "rule": rule})
            note = get(r, "Notes") if "Notes" in idx else ""
            if note:
                route["note"] = note
            check_rule(ws.title, sig, rule, aspect)

        for s in signals.values():
            for rt in s["routes"]:
                rt.pop("_key")
                order = {"clear": 0, "B2": 1, "B1": 2}
                rt["aspects"].sort(key=lambda a: order[a["when"]])
        cps.append({
            "id": m.group(1),
            "sheet": ws.title,
            "area": area,
            "switches": list(OrderedDict.fromkeys(h for _, h in sw_cols)),
            "signals": list(signals.values()),
        })

    cps.sort(key=lambda c: int(c["id"]))
    modified = wb.properties.modified.date().isoformat() if wb.properties.modified else ""
    return {
        "source": {"file": path.name, "version": version_of(path), "modified": modified},
        "controlPoints": cps,
    }


# C&O 1980 rule numbers (stock JMRI xml/signals/C&O-1980/aspects.xml), for a sanity check only.
CO_RULES = {
    "281": "Clear", "282": "Approach Medium", "283": "Medium Clear", "284": "Approach Slow",
    "285": "Approach", "286": "Medium Approach", "287": "Slow Clear", "288": "Slow Approach",
    "290": "Restricting", "291": "Stop and Proceed", "292": "Stop",
}
_rule_seen: set[tuple] = set()


def check_rule(sheet: str, sig: str, rule: str, aspect: str) -> None:
    want = CO_RULES.get(str(rule))
    norm = aspect.replace("&", "and").strip().lower()
    if want and want.lower() != norm and (rule, aspect) not in _rule_seen:
        _rule_seen.add((rule, aspect))
        warn(f"{sheet}: rule {rule} labeled '{aspect}' (C&O 1980 {rule} is '{want}'); first seen on {sig}")
    elif not want and rule and (rule, aspect) not in _rule_seen:
        _rule_seen.add((rule, aspect))
        warn(f"{sheet}: rule {rule} ('{aspect}') is not a C&O 1980 rule number; first seen on {sig}")


def prune(x):
    """Drop empty-string fields so the JSON only carries what a cell actually says."""
    if isinstance(x, dict):
        return {k: prune(v) for k, v in x.items() if v != ""}
    if isinstance(x, list):
        return [prune(v) for v in x]
    return x


# —— Main ——————————————————————————————————————————————————————————

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--inventory", type=Path, default=latest("LCOS_Layout_Inventory_v*.xlsx"),
                    help="LCOS Layout Inventory workbook (default: newest in docs/wiring/)")
    ap.add_argument("--splits", type=Path, default=latest("signals_split_v*.xlsx"),
                    help="signal split workbook (default: newest in docs/wiring/)")
    ap.add_argument("--out-dir", type=Path, default=PORTAL / "data", help="where the JSON goes")
    ap.add_argument("-q", "--quiet", action="store_true", help="don't list data warnings")
    args = ap.parse_args()
    if not args.inventory or not args.inventory.exists():
        ap.error(f"inventory workbook not found: {args.inventory}")
    if not args.splits or not args.splits.exists():
        ap.error(f"split workbook not found: {args.splits}")

    inv, ctx = build_inventory(args.inventory)
    splits = build_splits(args.splits, ctx)
    stamp = dt.date.today().isoformat()
    inv["source"]["built"] = stamp
    splits["source"]["built"] = stamp
    splits["source"]["inventory"] = inv["source"]["version"]

    args.out_dir.mkdir(parents=True, exist_ok=True)
    for name, data in (("wiring-inventory.json", inv), ("signal-splits.json", splits)):
        p = args.out_dir / name
        p.write_text(json.dumps(prune(data), indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"wrote {p} ({p.stat().st_size // 1024} KB)")

    n_pins = sum(len(b["pins"]) for n in inv["nodes"] for b in n["outputs"] + n["inputs"])
    print(f"{inv['source']['file']}: {len(inv['nodes'])} nodes, {n_pins} pins, {len(inv['signals'])} heads, "
          f"{len(inv['turnouts'])} turnouts, {len(inv['blocks'])} block sensors")
    print(f"{splits['source']['file']}: {len(splits['controlPoints'])} control points, "
          f"{sum(len(c['signals']) for c in splits['controlPoints'])} signals")
    if SCRUBBED:
        print(f"scrubbed {len(SCRUBBED)} cell(s) that looked like network details or secrets")
    if WARNINGS:
        print(f"{len(WARNINGS)} data warning(s)" + ("" if args.quiet else ":"))
        if not args.quiet:
            for w in WARNINGS:
                print("  - " + w)
    return 0


if __name__ == "__main__":
    sys.exit(main())
